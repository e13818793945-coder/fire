import os
import io
import json
from flask import Flask, request, session, redirect, url_for, render_template, flash, g, abort, send_file
from openpyxl import Workbook, load_workbook

import db as dbmod
from db import get_db, now_iso, today, get_setting, set_setting, log_action
from auth import load_logged_in_user, login_required, role_required, ROLE_LABEL, ROLE_HOME, all_agents
from kyc import generate_kyc_text, TIER_CHOICES, AGE_CHOICES, FAMILY_CHOICES, INCOME_CHOICES
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("COACHING_SECRET_KEY", "dev-secret-change-me")
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

dbmod.init_app(app)

NAV = {
    "admin": [("admin_users", "角色与权限管理"), ("admin_settings", "系统设置"), ("admin_audit_log", "操作日志")],
    "agent": [
        ("agent_clients", "我的客户"),
        ("agent_econ", "双周经营动态"),
        ("agent_feedback", "教练反馈"),
        ("agent_central", "线上辅导会总结"),
        ("agent_panke", "盘客报告"),
        ("agent_growth", "个人成长报告"),
    ],
    "coach": [("coach_updates", "经营动态反馈")],
    "pm": [
        ("pm_orphan", "孤儿单分配"),
        ("pm_periods", "经营动态周期"),
        ("pm_central", "集中辅导会录入"),
        ("pm_panke", "盘客报告录入"),
        ("pm_salon", "沙龙陪谈录入"),
        ("pm_growth", "成长报告录入"),
        ("pm_final", "结业报告录入"),
    ],
    "insurer": [
        ("insurer_overview", "经营动态总览"),
        ("insurer_central", "集中辅导会情况"),
        ("insurer_panke", "盘客辅导情况"),
        ("insurer_salon", "沙龙陪谈情况"),
        ("insurer_growth", "代理人成长报告"),
        ("insurer_final", "结业报告"),
    ],
}


@app.before_request
def _load_user():
    load_logged_in_user()


@app.context_processor
def _inject_nav():
    if g.get("user"):
        return {"nav_items": NAV.get(g.user["role"], []), "role_label": ROLE_LABEL.get(g.user["role"], "")}
    return {}


def render(template, **ctx):
    return render_template(template, **ctx)


# ---------------- auth ----------------
@app.route("/", methods=["GET"])
def index():
    if g.user is None:
        return redirect(url_for("login"))
    return redirect(url_for(ROLE_HOME[g.user["role"]]))


@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user is not None:
        return redirect(url_for(ROLE_HOME[g.user["role"]]))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        db = get_db()
        row = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if row is None or not check_password_hash(row["password_hash"], password):
            flash("账号或密码不正确，请重试", "error")
        elif not row["active"]:
            flash("该账号已被停用，请联系系统管理员", "error")
        else:
            session.clear()
            session["user_id"] = row["id"]
            log_action(row, "登录", "")
            return redirect(url_for(ROLE_HOME[row["role"]]))
    return render("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    if g.user is not None:
        log_action(g.user, "退出登录", "")
    session.clear()
    return redirect(url_for("login"))


# ================= 系统管理员 =================
@app.route("/admin/users")
@role_required("admin")
def admin_users():
    db = get_db()
    users = db.execute("SELECT * FROM users ORDER BY CASE role WHEN 'admin' THEN 0 WHEN 'pm' THEN 1 WHEN 'coach' THEN 2 WHEN 'agent' THEN 3 ELSE 4 END, display_name").fetchall()
    return render("admin/users.html", users=users, role_choices=[k for k in ROLE_LABEL if k != "admin"] + ["admin"], role_labels=ROLE_LABEL)


@app.route("/admin/users/new", methods=["POST"])
@role_required("admin")
def admin_users_new():
    db = get_db()
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    display_name = request.form.get("display_name", "").strip()
    role = request.form.get("role", "")
    if not username or not password or not display_name or role not in ROLE_LABEL:
        flash("请完整填写账号、密码、姓名并选择角色", "error")
        return redirect(url_for("admin_users"))
    exists = db.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone()
    if exists:
        flash(f"账号「{username}」已存在，请更换", "error")
        return redirect(url_for("admin_users"))
    lock_new = get_setting("lock_new_accounts", "0") == "1"
    db.execute(
        "INSERT INTO users (username, password_hash, role, display_name, active, created_at) VALUES (?,?,?,?,?,?)",
        (username, generate_password_hash(password), role, display_name, 0 if lock_new else 1, now_iso()),
    )
    db.commit()
    log_action(g.user, "创建账号", f"{username} / {ROLE_LABEL[role]}")
    flash(f"已创建账号「{username}」（{ROLE_LABEL[role]}）" + ("，账号默认锁定，需手动启用" if lock_new else ""), "ok")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@role_required("admin")
def admin_users_toggle(user_id):
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if row is None:
        abort(404)
    if row["role"] == "admin" and row["active"]:
        # 防止把自己或最后一个管理员锁死
        remaining = db.execute("SELECT COUNT(*) c FROM users WHERE role='admin' AND active=1 AND id!=?", (user_id,)).fetchone()["c"]
        if remaining == 0:
            flash("至少需要保留一个已启用的系统管理员账号", "error")
            return redirect(url_for("admin_users"))
    db.execute("UPDATE users SET active=? WHERE id=?", (0 if row["active"] else 1, user_id))
    db.commit()
    log_action(g.user, "启用/禁用账号", row["username"])
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/reset-password", methods=["POST"])
@role_required("admin")
def admin_users_reset_password(user_id):
    db = get_db()
    new_pass = request.form.get("new_password", "")
    row = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if row is None:
        abort(404)
    if not new_pass or len(new_pass) < 6:
        flash("新密码至少 6 位", "error")
        return redirect(url_for("admin_users"))
    db.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new_pass), user_id))
    db.commit()
    log_action(g.user, "重置密码", row["username"])
    flash(f"已重置「{row['username']}」的密码", "ok")
    return redirect(url_for("admin_users"))


@app.route("/admin/settings", methods=["GET"])
@role_required("admin")
def admin_settings():
    keys = ["auto_backup", "audit_log_retention_days", "lock_new_accounts"]
    settings = {k: get_setting(k) for k in keys}
    return render("admin/settings.html", settings=settings)


@app.route("/admin/settings/toggle/<key>", methods=["POST"])
@role_required("admin")
def admin_settings_toggle(key):
    if key not in ("auto_backup", "lock_new_accounts"):
        abort(404)
    current = get_setting(key, "0")
    set_setting(key, "0" if current == "1" else "1")
    log_action(g.user, "修改系统设置", f"{key} -> {'0' if current == '1' else '1'}")
    return redirect(url_for("admin_settings"))


@app.route("/admin/audit-log")
@role_required("admin")
def admin_audit_log():
    db = get_db()
    rows = db.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 200").fetchall()
    return render("admin/audit_log.html", rows=rows)


# ================= 代理人 =================
@app.route("/agent/clients")
@role_required("agent")
def agent_clients():
    db = get_db()
    clients = db.execute("SELECT * FROM clients WHERE agent_id=? ORDER BY created_at DESC", (g.user["id"],)).fetchall()
    return render("agent/clients.html", clients=clients)


@app.route("/agent/clients/new", methods=["POST"])
@role_required("agent")
def agent_clients_new():
    db = get_db()
    f = request.form
    name = f.get("name", "").strip()
    if not name:
        flash("请填写客户编码", "error")
        return redirect(url_for("agent_clients"))
    ts = now_iso()
    cur = db.execute(
        """INSERT INTO clients (agent_id, name, phone, tier, source, age_range, family_status, income_range, existing_policies, risk_notes, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        (g.user["id"], name, f.get("phone", ""), f.get("tier", "B"), "自有存量",
         f.get("age_range", ""), f.get("family_status", ""), f.get("income_range", ""),
         f.get("existing_policies", ""), f.get("risk_notes", ""), ts, ts),
    )
    client_id = cur.lastrowid
    client = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    kyc_text = generate_kyc_text(client)
    db.execute("INSERT INTO kyc_reports (client_id, content, generated_at) VALUES (?,?,?)", (client_id, kyc_text, now_iso()))
    db.commit()
    log_action(g.user, "录入客户", name)
    flash(f"已录入客户「{name}」，KYC 报告已自动生成", "ok")
    return redirect(url_for("agent_clients"))


@app.route("/agent/clients/<int:client_id>")
@role_required("agent")
def agent_client_detail(client_id):
    db = get_db()
    client = db.execute("SELECT * FROM clients WHERE id=? AND agent_id=?", (client_id, g.user["id"])).fetchone()
    if client is None:
        abort(404)
    kyc = db.execute("SELECT * FROM kyc_reports WHERE client_id=?", (client_id,)).fetchone()
    return render("agent/client_detail.html", client=client, kyc=kyc)


@app.route("/agent/econ", methods=["GET"])
@role_required("agent")
def agent_econ():
    db = get_db()
    periods = db.execute("SELECT * FROM econ_periods ORDER BY seq DESC").fetchall()
    current = db.execute("SELECT * FROM econ_periods WHERE is_current=1").fetchone()
    my_updates = {r["period_id"]: r for r in db.execute("SELECT * FROM econ_updates WHERE agent_id=?", (g.user["id"],)).fetchall()}
    feedback_map = {}
    if my_updates:
        rows = db.execute(
            f"SELECT * FROM coach_feedback WHERE econ_update_id IN ({','.join('?' * len(my_updates))})",
            tuple(u["id"] for u in my_updates.values()),
        ).fetchall()
        feedback_map = {r["econ_update_id"]: r for r in rows}
    return render("agent/econ.html", periods=periods, current=current, my_updates=my_updates, feedback_map=feedback_map)


@app.route("/agent/econ/submit", methods=["POST"])
@role_required("agent")
def agent_econ_submit():
    db = get_db()
    current = db.execute("SELECT * FROM econ_periods WHERE is_current=1").fetchone()
    if current is None:
        flash("项目经理尚未开启当前周期", "error")
        return redirect(url_for("agent_econ"))
    content = request.form.get("content", "").strip()
    if not content:
        flash("请填写本期经营动态", "error")
        return redirect(url_for("agent_econ"))
    existing = db.execute("SELECT * FROM econ_updates WHERE period_id=? AND agent_id=?", (current["id"], g.user["id"])).fetchone()
    if existing:
        db.execute("UPDATE econ_updates SET content=?, submitted_at=? WHERE id=?", (content, now_iso(), existing["id"]))
    else:
        db.execute(
            "INSERT INTO econ_updates (period_id, agent_id, content, submitted_at) VALUES (?,?,?,?)",
            (current["id"], g.user["id"], content, now_iso()),
        )
    db.commit()
    log_action(g.user, "提交经营动态", f"第{current['seq']}期")
    flash("本期经营动态已提交", "ok")
    return redirect(url_for("agent_econ"))


@app.route("/agent/feedback")
@role_required("agent")
def agent_feedback():
    db = get_db()
    rows = db.execute(
        """SELECT eu.*, ep.seq as period_seq, cf.content as fb_content, cf.submitted_at as fb_at, u.display_name as coach_name
           FROM econ_updates eu
           JOIN econ_periods ep ON ep.id = eu.period_id
           LEFT JOIN coach_feedback cf ON cf.econ_update_id = eu.id
           LEFT JOIN users u ON u.id = cf.coach_id
           WHERE eu.agent_id=? ORDER BY ep.seq DESC""",
        (g.user["id"],),
    ).fetchall()
    return render("agent/feedback.html", rows=rows)


@app.route("/agent/central")
@role_required("agent")
def agent_central():
    db = get_db()
    sessions = db.execute("SELECT * FROM central_sessions ORDER BY seq").fetchall()
    attendance = {r["session_id"]: r["present"] for r in db.execute("SELECT * FROM central_attendance WHERE agent_id=?", (g.user["id"],)).fetchall()}
    return render("agent/central.html", sessions=sessions, attendance=attendance)


@app.route("/agent/panke")
@role_required("agent")
def agent_panke():
    db = get_db()
    sessions = db.execute("SELECT * FROM panke_sessions ORDER BY seq").fetchall()
    attendance = {r["session_id"]: r["present"] for r in db.execute("SELECT * FROM panke_attendance WHERE agent_id=?", (g.user["id"],)).fetchall()}
    return render("agent/panke.html", sessions=sessions, attendance=attendance)


@app.route("/agent/growth")
@role_required("agent")
def agent_growth():
    db = get_db()
    report = db.execute("SELECT * FROM growth_reports WHERE agent_id=? AND status='published'", (g.user["id"],)).fetchone()
    metrics = json.loads(report["metrics_json"]) if report else {}
    return render("agent/growth.html", report=report, metrics=metrics)


# ================= 教练 =================
@app.route("/coach/updates")
@role_required("coach")
def coach_updates():
    db = get_db()
    current = db.execute("SELECT * FROM econ_periods WHERE is_current=1").fetchone()
    agents = all_agents(db)
    rows = []
    if current:
        updates = {r["agent_id"]: r for r in db.execute("SELECT * FROM econ_updates WHERE period_id=?", (current["id"],)).fetchall()}
        fb = {r["econ_update_id"]: r for r in db.execute(
            "SELECT * FROM coach_feedback WHERE econ_update_id IN (SELECT id FROM econ_updates WHERE period_id=?)", (current["id"],)
        ).fetchall()}
        for a in agents:
            u = updates.get(a["id"])
            rows.append({
                "agent": a,
                "econ_update": u,
                "feedback": fb.get(u["id"]) if u else None,
            })
    return render("coach/updates.html", current=current, rows=rows)


@app.route("/coach/updates/<int:update_id>/feedback", methods=["POST"])
@role_required("coach")
def coach_feedback_submit(update_id):
    db = get_db()
    u = db.execute("SELECT * FROM econ_updates WHERE id=?", (update_id,)).fetchone()
    if u is None:
        abort(404)
    content = request.form.get("content", "").strip()
    if not content:
        flash("请填写反馈内容", "error")
        return redirect(url_for("coach_updates"))
    existing = db.execute("SELECT * FROM coach_feedback WHERE econ_update_id=?", (update_id,)).fetchone()
    if existing:
        db.execute("UPDATE coach_feedback SET content=?, coach_id=?, submitted_at=? WHERE id=?", (content, g.user["id"], now_iso(), existing["id"]))
    else:
        db.execute(
            "INSERT INTO coach_feedback (econ_update_id, coach_id, content, submitted_at) VALUES (?,?,?,?)",
            (update_id, g.user["id"], content, now_iso()),
        )
    db.commit()
    log_action(g.user, "提交教练反馈", f"econ_update={update_id}")
    flash("反馈已提交", "ok")
    return redirect(url_for("coach_updates"))


# ================= 项目经理 =================
@app.route("/pm/orphan")
@role_required("pm")
def pm_orphan():
    db = get_db()
    agents = all_agents(db)
    orphan_clients = db.execute(
        "SELECT c.*, u.display_name as agent_name FROM clients c JOIN users u ON u.id=c.agent_id WHERE c.source='孤儿单' ORDER BY c.created_at DESC"
    ).fetchall()
    pool_rows = db.execute("SELECT * FROM orphan_pool ORDER BY imported_at DESC").fetchall()
    return render("pm/orphan.html", agents=agents, orphan_clients=orphan_clients, pool_rows=pool_rows)


@app.route("/pm/orphan/new", methods=["POST"])
@role_required("pm")
def pm_orphan_new():
    db = get_db()
    f = request.form
    name = f.get("name", "").strip()
    agent_id = f.get("agent_id", "")
    if not name or not agent_id:
        flash("请填写客户编码并选择分配的代理人", "error")
        return redirect(url_for("pm_orphan"))
    ts = now_iso()
    cur = db.execute(
        """INSERT INTO clients (agent_id, name, phone, tier, source, age_range, family_status, income_range, existing_policies, risk_notes, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        (agent_id, name, f.get("phone", ""), f.get("tier", "B"), "孤儿单", "", "", "", "", "", ts, ts),
    )
    client_id = cur.lastrowid
    client = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    db.execute("INSERT INTO kyc_reports (client_id, content, generated_at) VALUES (?,?,?)", (client_id, generate_kyc_text(client), now_iso()))
    db.commit()
    log_action(g.user, "分配孤儿单客户", f"{name} -> agent {agent_id}")
    flash(f"已将「{name}」分配给对应代理人", "ok")
    return redirect(url_for("pm_orphan"))


def _cell_str(v):
    """把 openpyxl 单元格的值统一转成去空格字符串；整数值的浮点数（如 20260901.0）去掉多余的 .0"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


@app.route("/pm/orphan/import/template")
@role_required("pm")
def pm_orphan_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "导入数据"
    ws.append(["客户编码", "分层(A/B/C)", "年龄区间", "家庭结构", "收入/资产区间"])

    note = wb.create_sheet("填写说明")
    note_lines = [
        "客户编码：必填，需保证在保司自己的名单里唯一；系统会自动跟已有编码去重。",
        "分层：A/B/C，留空默认按 B 处理。",
        "年龄区间可选值：" + "、".join(AGE_CHOICES) + "；无数据可留空。",
        "家庭结构可选值：" + "、".join(FAMILY_CHOICES) + "；无数据可留空。",
        "收入/资产区间可选值：" + "、".join(INCOME_CHOICES) + "；无数据可留空。",
        "留空的字段由代理人在分配后自行在系统内补充。",
        "重要：本模板任何一列都不要填写客户真实姓名、身份证号、手机号等可识别信息，只填写编码。",
    ]
    for i, line in enumerate(note_lines, start=1):
        note.cell(row=i, column=1, value=line)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="孤儿单导入模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/pm/orphan/import", methods=["POST"])
@role_required("pm")
def pm_orphan_import():
    file = request.files.get("import_file")
    if not file or not file.filename:
        flash("请选择要上传的文件", "error")
        return redirect(url_for("pm_orphan"))
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("仅支持 .xlsx 文件，请用上面的模板另存后上传", "error")
        return redirect(url_for("pm_orphan"))

    try:
        wb = load_workbook(file, data_only=True)
        ws = wb[wb.sheetnames[0]]
    except Exception:
        flash("文件解析失败，请确认是否为有效的 Excel 文件", "error")
        return redirect(url_for("pm_orphan"))

    db = get_db()
    existing_codes = {r["name"] for r in db.execute("SELECT name FROM clients").fetchall()}
    existing_codes |= {r["code"] for r in db.execute("SELECT code FROM orphan_pool").fetchall()}
    seen_in_file = set()

    ok_count = 0
    errors = []
    ts = now_iso()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or _cell_str(c) == "" for c in row):
            continue
        code = _cell_str(row[0]) if len(row) > 0 else ""
        tier = _cell_str(row[1]) if len(row) > 1 else ""
        age = _cell_str(row[2]) if len(row) > 2 else ""
        family = _cell_str(row[3]) if len(row) > 3 else ""
        income = _cell_str(row[4]) if len(row) > 4 else ""

        if not code:
            errors.append(f"第{idx}行：客户编码为空，已跳过")
            continue
        if code in existing_codes or code in seen_in_file:
            errors.append(f"第{idx}行：编码「{code}」重复，已跳过")
            continue
        if not tier:
            tier = "B"
        elif tier not in TIER_CHOICES:
            errors.append(f"第{idx}行「{code}」：分层「{tier}」不是 A/B/C，已跳过")
            continue
        if age and age not in AGE_CHOICES:
            errors.append(f"第{idx}行「{code}」：年龄区间「{age}」不在可选范围内，已跳过")
            continue
        if family and family not in FAMILY_CHOICES:
            errors.append(f"第{idx}行「{code}」：家庭结构「{family}」不在可选范围内，已跳过")
            continue
        if income and income not in INCOME_CHOICES:
            errors.append(f"第{idx}行「{code}」：收入区间「{income}」不在可选范围内，已跳过")
            continue

        db.execute(
            "INSERT INTO orphan_pool (code, tier, age_range, family_status, income_range, imported_by, imported_at) VALUES (?,?,?,?,?,?,?)",
            (code, tier, age, family, income, g.user["id"], ts),
        )
        seen_in_file.add(code)
        ok_count += 1

    db.commit()
    log_action(g.user, "批量导入孤儿单", f"成功{ok_count}条，失败{len(errors)}条")
    if ok_count:
        flash(f"成功导入 {ok_count} 条，进入待分配名单", "ok")
    if errors:
        shown = errors[:20]
        more = f"；另有 {len(errors) - 20} 条错误未列出" if len(errors) > 20 else ""
        flash("以下行未导入：" + "；".join(shown) + more, "error")
    if not ok_count and not errors:
        flash("文件中没有可导入的数据行", "error")
    return redirect(url_for("pm_orphan"))


@app.route("/pm/orphan/pool/<int:pool_id>/assign", methods=["POST"])
@role_required("pm")
def pm_orphan_pool_assign(pool_id):
    db = get_db()
    pool_row = db.execute("SELECT * FROM orphan_pool WHERE id=?", (pool_id,)).fetchone()
    if pool_row is None:
        abort(404)
    agent_id = request.form.get("agent_id", "")
    if not agent_id:
        flash("请选择分配的代理人", "error")
        return redirect(url_for("pm_orphan"))
    ts = now_iso()
    cur = db.execute(
        """INSERT INTO clients (agent_id, name, phone, tier, source, age_range, family_status, income_range, existing_policies, risk_notes, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        (agent_id, pool_row["code"], "", pool_row["tier"], "孤儿单",
         pool_row["age_range"] or "", pool_row["family_status"] or "", pool_row["income_range"] or "",
         "", "", ts, ts),
    )
    client_id = cur.lastrowid
    client = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    db.execute("INSERT INTO kyc_reports (client_id, content, generated_at) VALUES (?,?,?)", (client_id, generate_kyc_text(client), now_iso()))
    db.execute("DELETE FROM orphan_pool WHERE id=?", (pool_id,))
    db.commit()
    log_action(g.user, "分配孤儿单客户", f"{pool_row['code']} -> agent {agent_id}")
    flash(f"已将「{pool_row['code']}」分配给对应代理人", "ok")
    return redirect(url_for("pm_orphan"))


@app.route("/pm/orphan/pool/<int:pool_id>/delete", methods=["POST"])
@role_required("pm")
def pm_orphan_pool_delete(pool_id):
    db = get_db()
    row = db.execute("SELECT * FROM orphan_pool WHERE id=?", (pool_id,)).fetchone()
    if row is None:
        abort(404)
    db.execute("DELETE FROM orphan_pool WHERE id=?", (pool_id,))
    db.commit()
    log_action(g.user, "删除待分配孤儿单", row["code"])
    flash(f"已删除「{row['code']}」", "ok")
    return redirect(url_for("pm_orphan"))


@app.route("/pm/orphan/client/<int:client_id>/edit", methods=["POST"])
@role_required("pm")
def pm_orphan_client_edit(client_id):
    db = get_db()
    client = db.execute("SELECT * FROM clients WHERE id=? AND source='孤儿单'", (client_id,)).fetchone()
    if client is None:
        abort(404)
    f = request.form
    name = f.get("name", "").strip()
    tier = f.get("tier", "B")
    agent_id = f.get("agent_id", "")
    if not name or not agent_id:
        flash("请填写客户编码并选择代理人", "error")
        return redirect(url_for("pm_orphan"))
    if tier not in TIER_CHOICES:
        tier = "B"
    dup = db.execute("SELECT id FROM clients WHERE name=? AND id!=?", (name, client_id)).fetchone()
    if dup:
        flash(f"客户编码「{name}」已被其他客户占用，未保存", "error")
        return redirect(url_for("pm_orphan"))
    db.execute(
        "UPDATE clients SET name=?, tier=?, agent_id=?, updated_at=? WHERE id=?",
        (name, tier, agent_id, now_iso(), client_id),
    )
    updated = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    db.execute("DELETE FROM kyc_reports WHERE client_id=?", (client_id,))
    db.execute("INSERT INTO kyc_reports (client_id, content, generated_at) VALUES (?,?,?)", (client_id, generate_kyc_text(updated), now_iso()))
    db.commit()
    log_action(g.user, "编辑孤儿单客户", f"{client['name']} -> {name}")
    flash(f"「{name}」已更新，KYC 报告已重新生成", "ok")
    return redirect(url_for("pm_orphan"))


@app.route("/pm/orphan/client/<int:client_id>/delete", methods=["POST"])
@role_required("pm")
def pm_orphan_client_delete(client_id):
    db = get_db()
    client = db.execute("SELECT * FROM clients WHERE id=? AND source='孤儿单'", (client_id,)).fetchone()
    if client is None:
        abort(404)
    db.execute("DELETE FROM kyc_reports WHERE client_id=?", (client_id,))
    db.execute("DELETE FROM clients WHERE id=?", (client_id,))
    db.commit()
    log_action(g.user, "删除孤儿单客户", client["name"])
    flash(f"已删除客户「{client['name']}」及其 KYC 报告", "ok")
    return redirect(url_for("pm_orphan"))


@app.route("/pm/periods")
@role_required("pm")
def pm_periods():
    db = get_db()
    periods = db.execute("SELECT * FROM econ_periods ORDER BY seq DESC").fetchall()
    return render("pm/periods.html", periods=periods)


@app.route("/pm/periods/new", methods=["POST"])
@role_required("pm")
def pm_periods_new():
    db = get_db()
    start = request.form.get("start_date", "")
    end = request.form.get("end_date", "")
    if not start or not end:
        flash("请填写周期起止日期", "error")
        return redirect(url_for("pm_periods"))
    last = db.execute("SELECT MAX(seq) m FROM econ_periods").fetchone()["m"] or 0
    db.execute("UPDATE econ_periods SET is_current=0")
    db.execute("INSERT INTO econ_periods (seq, start_date, end_date, is_current) VALUES (?,?,?,1)", (last + 1, start, end))
    db.commit()
    log_action(g.user, "新增经营动态周期", f"第{last+1}期 {start}~{end}")
    flash(f"已开启第 {last + 1} 期，并设为当前周期", "ok")
    return redirect(url_for("pm_periods"))


@app.route("/pm/periods/<int:period_id>/set-current", methods=["POST"])
@role_required("pm")
def pm_periods_set_current(period_id):
    db = get_db()
    db.execute("UPDATE econ_periods SET is_current=0")
    db.execute("UPDATE econ_periods SET is_current=1 WHERE id=?", (period_id,))
    db.commit()
    log_action(g.user, "切换当前周期", str(period_id))
    return redirect(url_for("pm_periods"))


def _session_view(kind):
    db = get_db()
    table = "central_sessions" if kind == "central" else "panke_sessions"
    att_table = "central_attendance" if kind == "central" else "panke_attendance"
    sessions = db.execute(f"SELECT * FROM {table} ORDER BY seq").fetchall()
    agents = all_agents(db)
    att = {}
    for r in db.execute(f"SELECT * FROM {att_table}").fetchall():
        att[(r["session_id"], r["agent_id"])] = r["present"]
    return sessions, agents, att


@app.route("/pm/central")
@role_required("pm")
def pm_central():
    sessions, agents, att = _session_view("central")
    return render("pm/central.html", sessions=sessions, agents=agents, att=att)


@app.route("/pm/central/<int:session_id>/update", methods=["POST"])
@role_required("pm")
def pm_central_update(session_id):
    db = get_db()
    f = request.form
    db.execute(
        "UPDATE central_sessions SET date=?, title=?, summary=?, published=? WHERE id=?",
        (f.get("date", ""), f.get("title", ""), f.get("summary", ""), 1 if f.get("published") == "on" else 0, session_id),
    )
    db.commit()
    log_action(g.user, "更新集中辅导会总结", f"session={session_id}")
    flash("已保存", "ok")
    return redirect(url_for("pm_central"))


@app.route("/pm/central/<int:session_id>/attendance/<int:agent_id>/toggle", methods=["POST"])
@role_required("pm")
def pm_central_attendance_toggle(session_id, agent_id):
    db = get_db()
    row = db.execute("SELECT * FROM central_attendance WHERE session_id=? AND agent_id=?", (session_id, agent_id)).fetchone()
    if row:
        db.execute("UPDATE central_attendance SET present=? WHERE id=?", (0 if row["present"] else 1, row["id"]))
    else:
        db.execute("INSERT INTO central_attendance (session_id, agent_id, present) VALUES (?,?,1)", (session_id, agent_id))
    db.commit()
    return redirect(url_for("pm_central"))


@app.route("/pm/panke")
@role_required("pm")
def pm_panke():
    sessions, agents, att = _session_view("panke")
    return render("pm/panke.html", sessions=sessions, agents=agents, att=att)


@app.route("/pm/panke/<int:session_id>/update", methods=["POST"])
@role_required("pm")
def pm_panke_update(session_id):
    db = get_db()
    f = request.form
    db.execute(
        "UPDATE panke_sessions SET date=?, focus=?, content=?, published=? WHERE id=?",
        (f.get("date", ""), f.get("focus", ""), f.get("content", ""), 1 if f.get("published") == "on" else 0, session_id),
    )
    db.commit()
    log_action(g.user, "更新盘客报告", f"session={session_id}")
    flash("已保存", "ok")
    return redirect(url_for("pm_panke"))


@app.route("/pm/panke/<int:session_id>/attendance/<int:agent_id>/toggle", methods=["POST"])
@role_required("pm")
def pm_panke_attendance_toggle(session_id, agent_id):
    db = get_db()
    row = db.execute("SELECT * FROM panke_attendance WHERE session_id=? AND agent_id=?", (session_id, agent_id)).fetchone()
    if row:
        db.execute("UPDATE panke_attendance SET present=? WHERE id=?", (0 if row["present"] else 1, row["id"]))
    else:
        db.execute("INSERT INTO panke_attendance (session_id, agent_id, present) VALUES (?,?,1)", (session_id, agent_id))
    db.commit()
    return redirect(url_for("pm_panke"))


@app.route("/pm/salon")
@role_required("pm")
def pm_salon():
    db = get_db()
    sessions = db.execute("SELECT * FROM salon_sessions ORDER BY seq").fetchall()
    agents = all_agents(db)
    notes = {(r["session_id"], r["agent_id"]): r for r in db.execute("SELECT * FROM salon_notes").fetchall()}
    return render("pm/salon.html", sessions=sessions, agents=agents, notes=notes)


@app.route("/pm/salon/<int:session_id>/save", methods=["POST"])
@role_required("pm")
def pm_salon_save(session_id):
    db = get_db()
    agents = all_agents(db)
    for a in agents:
        note = request.form.get(f"note_{a['id']}", "").strip()
        present = 1 if request.form.get(f"present_{a['id']}") == "on" else 0
        row = db.execute("SELECT * FROM salon_notes WHERE session_id=? AND agent_id=?", (session_id, a["id"])).fetchone()
        if row:
            db.execute("UPDATE salon_notes SET note=?, present=? WHERE id=?", (note, present, row["id"]))
        else:
            db.execute("INSERT INTO salon_notes (session_id, agent_id, note, present) VALUES (?,?,?,?)", (session_id, a["id"], note, present))
    db.commit()
    log_action(g.user, "更新沙龙陪谈记录", f"session={session_id}")
    flash("已保存", "ok")
    return redirect(url_for("pm_salon"))


@app.route("/pm/growth")
@role_required("pm")
def pm_growth():
    db = get_db()
    agents = all_agents(db)
    reports = {r["agent_id"]: r for r in db.execute("SELECT * FROM growth_reports").fetchall()}
    return render("pm/growth.html", agents=agents, reports=reports)


@app.route("/pm/growth/<int:agent_id>/edit")
@role_required("pm")
def pm_growth_edit(agent_id):
    db = get_db()
    agent = db.execute("SELECT * FROM users WHERE id=? AND role='agent'", (agent_id,)).fetchone()
    if agent is None:
        abort(404)
    report = db.execute("SELECT * FROM growth_reports WHERE agent_id=?", (agent_id,)).fetchone()
    metrics = json.loads(report["metrics_json"]) if report else {}
    dims = ["客户经营能力", "KYC应用能力", "活动量达成", "转介绍开发"]
    return render("pm/growth_edit.html", agent=agent, report=report, metrics=metrics, dims=dims)


@app.route("/pm/growth/<int:agent_id>/save", methods=["POST"])
@role_required("pm")
def pm_growth_save(agent_id):
    db = get_db()
    dims = ["客户经营能力", "KYC应用能力", "活动量达成", "转介绍开发"]
    metrics = {}
    for d in dims:
        try:
            metrics[d] = max(0, min(100, int(request.form.get(f"metric_{d}", "0"))))
        except ValueError:
            metrics[d] = 0
    narrative = request.form.get("narrative", "").strip()
    status = "published" if request.form.get("published") == "on" else "draft"
    row = db.execute("SELECT * FROM growth_reports WHERE agent_id=?", (agent_id,)).fetchone()
    if row:
        db.execute(
            "UPDATE growth_reports SET metrics_json=?, narrative=?, status=?, updated_at=? WHERE id=?",
            (json.dumps(metrics, ensure_ascii=False), narrative, status, now_iso(), row["id"]),
        )
    else:
        db.execute(
            "INSERT INTO growth_reports (agent_id, metrics_json, narrative, status, updated_at) VALUES (?,?,?,?,?)",
            (agent_id, json.dumps(metrics, ensure_ascii=False), narrative, status, now_iso()),
        )
    db.commit()
    log_action(g.user, "更新成长报告", f"agent={agent_id} status={status}")
    flash("成长报告已保存", "ok")
    return redirect(url_for("pm_growth"))


@app.route("/pm/final")
@role_required("pm")
def pm_final():
    db = get_db()
    report = db.execute("SELECT * FROM final_report WHERE id=1").fetchone()
    return render("pm/final.html", report=report)


@app.route("/pm/final/save", methods=["POST"])
@role_required("pm")
def pm_final_save():
    db = get_db()
    content = request.form.get("content", "").strip()
    published = 1 if request.form.get("published") == "on" else 0
    db.execute("UPDATE final_report SET content=?, published=?, updated_at=? WHERE id=1", (content, published, now_iso()))
    db.commit()
    log_action(g.user, "更新结业报告", f"published={published}")
    flash("结业报告已保存", "ok")
    return redirect(url_for("pm_final"))


# ================= 保司内勤（只读） =================
@app.route("/insurer/overview")
@role_required("insurer")
def insurer_overview():
    db = get_db()
    current = db.execute("SELECT * FROM econ_periods WHERE is_current=1").fetchone()
    agents = all_agents(db)
    rows = []
    if current:
        updates = {r["agent_id"]: r for r in db.execute("SELECT * FROM econ_updates WHERE period_id=?", (current["id"],)).fetchall()}
        fb_ids = {u["id"] for u in updates.values()}
        fb = {}
        if fb_ids:
            fb = {r["econ_update_id"]: r for r in db.execute(
                f"SELECT * FROM coach_feedback WHERE econ_update_id IN ({','.join('?' * len(fb_ids))})", tuple(fb_ids)
            ).fetchall()}
        for a in agents:
            u = updates.get(a["id"])
            rows.append({"agent": a, "submitted": bool(u), "feedback": bool(u and u["id"] in fb)})
    return render("insurer/overview.html", current=current, rows=rows)


@app.route("/insurer/central")
@role_required("insurer")
def insurer_central():
    db = get_db()
    sessions = db.execute("SELECT * FROM central_sessions ORDER BY seq").fetchall()
    agents = all_agents(db)
    att = {(r["session_id"], r["agent_id"]): r["present"] for r in db.execute("SELECT * FROM central_attendance").fetchall()}
    return render("insurer/central.html", sessions=sessions, agents=agents, att=att)


@app.route("/insurer/panke")
@role_required("insurer")
def insurer_panke():
    db = get_db()
    sessions = db.execute("SELECT * FROM panke_sessions ORDER BY seq").fetchall()
    agents = all_agents(db)
    att = {(r["session_id"], r["agent_id"]): r["present"] for r in db.execute("SELECT * FROM panke_attendance").fetchall()}
    return render("insurer/panke.html", sessions=sessions, agents=agents, att=att)


@app.route("/insurer/salon")
@role_required("insurer")
def insurer_salon():
    db = get_db()
    sessions = db.execute("SELECT * FROM salon_sessions ORDER BY seq").fetchall()
    agents = all_agents(db)
    notes = {(r["session_id"], r["agent_id"]): r for r in db.execute("SELECT * FROM salon_notes").fetchall()}
    return render("insurer/salon.html", sessions=sessions, agents=agents, notes=notes)


@app.route("/insurer/growth")
@role_required("insurer")
def insurer_growth():
    db = get_db()
    agents = all_agents(db)
    reports = {r["agent_id"]: r for r in db.execute("SELECT * FROM growth_reports WHERE status='published'").fetchall()}
    parsed = {aid: json.loads(r["metrics_json"]) for aid, r in reports.items()}
    return render("insurer/growth.html", agents=agents, reports=reports, metrics=parsed)


@app.route("/insurer/final")
@role_required("insurer")
def insurer_final():
    db = get_db()
    report = db.execute("SELECT * FROM final_report WHERE id=1").fetchone()
    return render("insurer/final.html", report=report)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=os.environ.get("FLASK_DEBUG") == "1")
